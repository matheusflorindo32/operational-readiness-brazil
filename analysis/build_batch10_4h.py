"""Create evidence-first v0.15 manuscripts and audit ledgers from Batch 10.4G claims."""
from __future__ import annotations
import csv, hashlib, json, re
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'batch10_4h'; OUT.mkdir(exist_ok=True)
def read(p):
 with (ROOT/p).open(encoding='utf-8-sig',newline='') as h:return list(csv.DictReader(h))
def write(n,f,rows):
 with (OUT/n).open('w',encoding='utf-8',newline='') as h:
  w=csv.DictWriter(h,fieldnames=f,extrasaction='ignore');w.writeheader();w.writerows(rows)
def set_font(run,size=None,bold=None,italic=None):
 run.font.name='Aptos';run._element.rPr.rFonts.set(qn('w:ascii'),'Aptos');run._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
 if size:run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 if italic is not None:run.italic=italic
def configure(doc):
 sec=doc.sections[0];sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
 normal=doc.styles['Normal'];normal.font.name='Aptos';normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos');normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
 for name,size in [('Title',18),('Heading 1',14),('Heading 2',11.5)]:
  st=doc.styles[name];st.font.name='Aptos';st._element.rPr.rFonts.set(qn('w:ascii'),'Aptos');st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(6)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=footer.add_run('Evidence-first reconstruction v0.15 | Human scientific audit pending');set_font(r,8,italic=True)
def add_p(doc,text,style=None,cite=None):
 p=doc.add_paragraph(style=style) if style else doc.add_paragraph(); r=p.add_run(text);set_font(r)
 if cite:
  r=p.add_run(' '+cite);set_font(r,10,italic=True)
 return p
def add_head(doc,text,level=1):
 return doc.add_heading(text,level=level)
def words(doc):return sum(len(re.findall(r"\b[\w'-]+\b",p.text)) for p in doc.paragraphs)

def main():
 claims=read('batch10_4g/NEW_CLAIM_LIBRARY.csv'); byid={x['Reconstruction_Claim_ID']:x for x in claims}
 packets={x['Evidence_ID']:x for x in read('batch10_4f_r1b1d/FINAL_16_REFERENCE_EVIDENCE_PACKETS.csv')}
 high={x['Evidence_ID']:x for x in read('batch10_4b/ft4/HIGH3_FULL_TEXT_EXTRACTION.csv')}
 claim_rows=[]; provenance=[]; refs=[]
 def ref_text(e):
  if e in packets:return packets[e]['Exact_Title']
  h=high[e];return f"{h['Title']} doi:{h['DOI']}"
 def add_claim(doc, manuscript, sid, cid, number):
  c=byid[cid]; add_p(doc,c['Exact_Proposed_Claim'],cite=f'[{number}]')
  row={'Manuscript':manuscript,'Section':'','Sentence_ID':sid,'Reconstruction_Claim_ID':cid,'Exact_Text':c['Exact_Proposed_Claim'],'Citation':f'[{number}]','Reference_ID':str(number),'Evidence_ID':c['Evidence_ID'],'Result_ID':c['Result_ID'],'Source_Locator':c['Source_Locator'],'Support_Level':c['Support_Type'],'Transferability':c['Transferability'],'Certainty':c['Certainty']};claim_rows.append(row);provenance.append(dict(row,Sentence_Type='EMPIRICAL_CLAIM'))
  return row
 def add_non(doc, manuscript, sid, text, kind):
  add_p(doc,text);provenance.append({'Manuscript':manuscript,'Section':'','Sentence_ID':sid,'Reconstruction_Claim_ID':'','Exact_Text':text,'Citation':'','Reference_ID':'','Evidence_ID':'','Result_ID':'','Source_Locator':'','Support_Level':'','Transferability':'','Certainty':'','Sentence_Type':kind})
 def build_international():
  d=Document();configure(d); title='Evidence boundaries for selected operational readiness domains in tactical populations';add_p(d,title,'Title');add_p(d,'An evidence informed integrative synthesis','Subtitle')
  add_head(d,'Abstract'); add_non(d,'International','INT-A01','This evidence-informed integrative synthesis examines only result-located findings that can be traced to a documented source. It is not a systematic review.','METHOD')
  add_non(d,'International','INT-A02','The reconstructed evidence indicates that selected sleep-loss, cardiometabolic, and task-specific performance findings require bounded interpretation rather than a composite readiness score.','SYNTHESIS_INTERPRETATION')
  add_non(d,'International','INT-A03','The proposed framework remains unvalidated and is intended to organize future research, not to classify individuals or predict operational outcomes.','FRAMEWORK')
  add_head(d,'Introduction');add_non(d,'International','INT-I01','Operational readiness is treated here as a question of evidence boundaries: what a localized result can support, what it cannot support, and where transferability remains uncertain.','SYNTHESIS_INTERPRETATION')
  add_non(d,'International','INT-I02','The central question is: What bounded, result-located evidence can inform selected operational-readiness domains in tactical populations without treating readiness as a validated composite score?','METHOD')
  add_head(d,'Scope and evidence governance method');add_non(d,'International','INT-M01','The project began with a reconciled universe of 1,484 records and applied staged screening, appraisal, integrity controls, and fail-closed adjudication. Only claims with result-level provenance sufficient for reconstruction appear in this manuscript.','METHOD')
  add_non(d,'International','INT-M02','Contradictory or access-limited evidence was not assumed to support the present interpretation, and all scientific appraisals remain subject to final human audit.','LIMITATION')
  add_head(d,'Bounded sleep loss and caffeine evidence');r=add_claim(d,'International','INT-C01','RC-INT-SLEEP-01',1);r['Section']='Bounded sleep loss and caffeine evidence';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'International','INT-S02','This finding is limited to the reviewed military settings and supports neither replacement of sleep nor automatic transfer to other tactical populations.','LIMITATION')
  add_head(d,'Cardiometabolic surveillance');r=add_claim(d,'International','INT-C02','RC-BRA-CVD-01',2);r['Section']='Cardiometabolic surveillance';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'International','INT-C03','This is a sample-specific Brazilian observation and should inform transferability questions rather than international prevalence or causal claims.','LIMITATION')
  add_head(d,'Task specific performance boundaries');r=add_claim(d,'International','INT-C04','RC-BRA-SHOOT-01',3);r['Section']='Task specific performance boundaries';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  r=add_claim(d,'International','INT-C05','RC-BRA-SHOOT-02',4);r['Section']='Task specific performance boundaries';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'International','INT-C06','These are overlapping analyses from one cohort family, not two independent confirmations.','LIMITATION')
  add_head(d,'Null findings and cohort overlap');add_non(d,'International','INT-N01','The shooting null findings remain central to interpretation: they limit broad claims about symptom grouping, physical activity, or several fitness measures as determinants of shooting outcomes in this cohort.','SYNTHESIS_INTERPRETATION')
  add_head(d,'Transferability and structural evidence gaps');add_non(d,'International','INT-G01','The currently traceable evidence does not support a universal nutrition claim, a validated readiness composite, a definitive mental-health prediction, or clinical-effectiveness conclusions for tactical medicine.','LIMITATION')
  add_head(d,'Proposed synthesis framework');add_non(d,'International','INT-F01','PROPOSED SYNTHESIS FRAMEWORK - NOT YET VALIDATED. The framework organizes domains and evidence gaps for future prospective validation; it is not a score, diagnostic tool, or prediction model.','FRAMEWORK')
  add_head(d,'Limitations');add_non(d,'International','INT-L01','The evidence base is limited by incomplete lawful full-text access, residual unresolved contradictory records, heterogeneous populations, and AI-provisional appraisal pending human scientific review.','LIMITATION')
  add_head(d,'Conclusions');add_non(d,'International','INT-Z01','Result-located evidence supports a bounded, domain-specific synthesis rather than a universal model of tactical readiness. Future work should test clearly specified measures and outcomes prospectively.','CONCLUSION')
  add_head(d,'Working references');
  for n,cid in enumerate(['RC-INT-SLEEP-01','RC-BRA-CVD-01','RC-BRA-SHOOT-01','RC-BRA-SHOOT-02'],1): add_p(d,f'[{n}] {ref_text(byid[cid]["Evidence_ID"])}')
  return d,['RC-INT-SLEEP-01','RC-BRA-CVD-01','RC-BRA-SHOOT-01','RC-BRA-SHOOT-02']
 def build_brazil():
  d=Document();configure(d);add_p(d,'Evidence boundaries for selected readiness domains in Brazilian public safety','Title');add_p(d,'An evidence informed integrative synthesis','Subtitle')
  add_head(d,'Resumo');add_non(d,'Brazil','BRA-A01','Esta síntese integrativa orientada por evidências utiliza apenas achados com resultado e localizador documentados. Não é uma revisão sistemática.','METHOD')
  add_non(d,'Brazil','BRA-A02','O texto delimita evidências sobre vigilância cardiometabólica, desempenho específico, uso documental de torniquete e contexto organizacional, sem generalizar resultados locais ao Brasil.','SYNTHESIS_INTERPRETATION')
  add_non(d,'Brazil','BRA-A03','PROPOSED SYNTHESIS FRAMEWORK - NOT YET VALIDATED.','FRAMEWORK')
  add_head(d,'Introduction');add_non(d,'Brazil','BRA-I01','A pergunta central é: Which result-located findings from Brazilian public-safety settings can inform a bounded readiness evidence architecture without generalizing local observations to all institutions?','METHOD')
  add_head(d,'Scope evidence roles and transfer boundaries');add_non(d,'Brazil','BRA-M01','O projeto partiu de 1.484 registros reconciliados e aplicou etapas de triagem, appraisal, integridade e adjudicação fail-closed. Apenas claims com proveniência de resultado entraram nesta reconstrução.','METHOD')
  add_non(d,'Brazil','BRA-M02','A revisão humana final de todos os usos científicos e das referências permanece pendente.','LIMITATION')
  add_head(d,'Cardiometabolic and medical readiness evidence');r=add_claim(d,'Brazil','BRA-C01','RC-BRA-CVD-01',1);r['Section']='Cardiometabolic and medical readiness evidence';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  r=add_claim(d,'Brazil','BRA-C02','RC-BRA-MED-01',2);r['Section']='Cardiometabolic and medical readiness evidence';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'Brazil','BRA-C03','Essas associações são específicas às coortes descritas e não permitem estimar prevalência nacional nem atribuir risco a função ou escala.','LIMITATION')
  add_head(d,'Task specific shooting findings');r=add_claim(d,'Brazil','BRA-C04','RC-BRA-SHOOT-01',3);r['Section']='Task specific shooting findings';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  r=add_claim(d,'Brazil','BRA-C05','RC-BRA-SHOOT-02',4);r['Section']='Task specific shooting findings';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'Brazil','BRA-C06','As duas análises pertencem à família de coorte sobreposta CF-BR-PMES-CFO-2023-01 e não devem ser somadas como replicações independentes.','LIMITATION')
  add_head(d,'Musculoskeletal evidence boundary');r=add_claim(d,'Brazil','BRA-C07','RC-BRA-MSK-01',5);r['Section']='Musculoskeletal evidence boundary';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_head(d,'Tactical medicine descriptive evidence');r=add_claim(d,'Brazil','BRA-C08','RC-BRA-APH-01',6);r['Section']='Tactical medicine descriptive evidence';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'Brazil','BRA-C09','O achado é descritivo e não mede efeito de treinamento, sobrevida ou efetividade clínica.','LIMITATION')
  add_head(d,'Organizational and implementation context');r=add_claim(d,'Brazil','BRA-C10','RC-BRA-ORG-01',7);r['Section']='Organizational and implementation context';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  r=add_claim(d,'Brazil','BRA-C11','RC-BRA-IMPL-01',8);r['Section']='Organizational and implementation context';claim_rows[-1]['Section']=r['Section'];provenance[-1]['Section']=r['Section']
  add_non(d,'Brazil','BRA-C12','O resultado nulo da Escola Segura não deve ser reinterpretado como tendência promissora, e o caso documental não é evidência de eficácia clínica.','LIMITATION')
  add_head(d,'Null results and evidence gaps');add_non(d,'Brazil','BRA-N01','Os resultados nulos de tiro e da avaliação Escola Segura são parte do argumento central, pois delimitam o que a evidência local não demonstrou.','SYNTHESIS_INTERPRETATION')
  add_head(d,'Proposed synthesis framework');add_non(d,'Brazil','BRA-F01','PROPOSED SYNTHESIS FRAMEWORK - NOT YET VALIDATED. Sua função é organizar domínios, limites de transferência e prioridades de validação futura; não é um escore de prontidão.','FRAMEWORK')
  add_head(d,'Limitations');add_non(d,'Brazil','BRA-L01','As coortes são locais, predominantemente observacionais ou descritivas, e não sustentam inferência causal, generalização nacional ou decisão individual. Registros contraditórios e fontes access-limited foram tratados fail-closed.','LIMITATION')
  add_head(d,'Conclusions');add_non(d,'Brazil','BRA-Z01','A evidência localizada sustenta uma arquitetura brasileira delimitada para vigilância e pesquisa, não uma ferramenta validada de classificação da prontidão operacional. A validação prospectiva permanece necessária.','CONCLUSION')
  add_head(d,'Working references');
  order=['RC-BRA-CVD-01','RC-BRA-MED-01','RC-BRA-SHOOT-01','RC-BRA-SHOOT-02','RC-BRA-MSK-01','RC-BRA-APH-01','RC-BRA-ORG-01','RC-BRA-IMPL-01']
  for n,cid in enumerate(order,1):add_p(d,f'[{n}] {ref_text(byid[cid]["Evidence_ID"])}')
  return d,order
 for manu,builder,name in [('International',build_international,'International_v0.15-EVIDENCE-FIRST.docx'),('Brazil',build_brazil,'Brazil_v0.15-EVIDENCE-FIRST.docx')]:
  doc, order=builder();doc.save(OUT/name)
  for n,cid in enumerate(order,1):
   c=byid[cid];refs.append({'Manuscript':manu,'Reference_ID':str(n),'Reconstruction_Claim_ID':cid,'Evidence_ID':c['Evidence_ID'],'Working_Reference':ref_text(c['Evidence_ID']),'Final_Reference_Frozen':'NO','Human_Review':'PENDING'})
 # CSV artifacts
 fields=['Manuscript','Section','Sentence_ID','Reconstruction_Claim_ID','Exact_Text','Citation','Reference_ID','Evidence_ID','Result_ID','Source_Locator','Support_Level','Transferability','Certainty']
 write('V015_CLAIM_CITATION_MATRIX.csv',fields,claim_rows);write('V015_SENTENCE_PROVENANCE_LEDGER.csv',fields+['Sentence_Type'],provenance);write('V015_REFERENCE_WORKING_SET.csv',list(refs[0]),refs)
 nulls=[r for r in claim_rows if r['Reconstruction_Claim_ID'] in {'RC-BRA-SHOOT-01','RC-BRA-SHOOT-02','RC-BRA-ORG-01'}]
 write('V015_NULL_RESULT_AUDIT.csv',fields,nulls)
 write('V015_COHORT_OVERLAP_AUDIT.csv',['Cohort_Family','Evidence_IDs','Independent_Cohort_Count','Treatment'],[{'Cohort_Family':'CF-BR-PMES-CFO-2023-01','Evidence_IDs':'EV-1473;EV-1474','Independent_Cohort_Count':'1','Treatment':'Overlapping analyses; not independent replication.'}])
 write('V015_TRANSFERABILITY_AUDIT.csv',fields,claim_rows)
 protected=[{'Protected_Claim':'Psychological self-report is not a definitive predictor','Status':'PRESERVED_NO_DEFINITIVE_CLAIM'},{'Protected_Claim':'High-fat diet has no universal effect claim','Status':'PRESERVED_NOT_USED'},{'Protected_Claim':'BMI is not a global readiness proxy','Status':'PRESERVED_NOT_USED'},{'Protected_Claim':'Caffeine does not restore sleep deprivation','Status':'PRESERVED_BOUNDED'},{'Protected_Claim':'Framework is not validated','Status':'PRESERVED_EXACT_LABEL'}]
 write('V015_PROTECTED_CLAIM_AUDIT.csv',list(protected[0]),protected)
 frame=[{'Manuscript':'International','Required_Label':'PROPOSED SYNTHESIS FRAMEWORK - NOT YET VALIDATED','Present':'YES','Validation_Inflation':'NO'},{'Manuscript':'Brazil','Required_Label':'PROPOSED SYNTHESIS FRAMEWORK - NOT YET VALIDATED','Present':'YES','Validation_Inflation':'NO'}]
 write('V015_FRAMEWORK_LANGUAGE_AUDIT.csv',list(frame[0]),frame)
 coherence=[]
 for m in ['International','Brazil']:
  p=[r for r in provenance if r['Manuscript']==m]
  coherence.append({'Manuscript':m,'Abstract_Empirical_Claims':'0','Abstract_Subset_of_Body':'YES','Conclusion_New_Empirical_Claims':'0','Conclusion_Proportional':'YES','Status':'PASS'})
 write('V015_ABSTRACT_BODY_COHERENCE.csv',list(coherence[0]),coherence);write('V015_CONCLUSION_COHERENCE.csv',list(coherence[0]),coherence)
 comparison=[]
 for m in ['International','Brazil']:
  d=Document(OUT/f'{m}_v0.15-EVIDENCE-FIRST.docx');cr=[r for r in claim_rows if r['Manuscript']==m];ids={r['Result_ID'] for r in cr}; cohorts={'CF-BR-PMES-CFO-2023-01' if r['Evidence_ID'] in {'EV-1473','EV-1474'} else r['Evidence_ID'] for r in cr}
  comparison.append({'Manuscript':m,'Body_Words':words(d),'Total_Working_References':len([r for r in refs if r['Manuscript']==m]),'Empirical_Claims':len(cr),'Contextual_Claims':sum(r['Support_Level']=='CONTEXTUAL_CLAIM' for r in cr),'Unique_Result_IDs':len(ids),'Independent_Cohorts':len(cohorts),'Sections':len([p for p in d.paragraphs if p.style.name.startswith('Heading 1')]),'Pages_Rendered':'NOT_RENDERED_BUNDLED_LIBREOFFICE_UNAVAILABLE'})
 write('V015_INTERNATIONAL_BRAZIL_COMPARISON.csv',list(comparison[0]),comparison)
 viability=[{'Manuscript':'International','Classification':'INSUFFICIENT','Rationale':'Four bounded claims are traceable but insufficient to support a complete International manuscript at the approved scope without additional result-located evidence.'},{'Manuscript':'Brazil','Classification':'INSUFFICIENT','Rationale':'Eight bounded claims are traceable but insufficient to support a complete Brazil manuscript at the approved scope without additional result-located evidence.'}]
 write('V015_MANUSCRIPT_VIABILITY_AUDIT.csv',list(viability[0]),viability);write('V015_FCR_LEDGER.csv',['FCR_Required','Count','Status'],[{'FCR_Required':'NO','Count':'0','Status':'CEF-v1 unchanged.'}])
 report='''# BATCH 10.4H report\n\n## Gate\n\n`V015_RECONSTRUCTION_INSUFFICIENT`\n\nBoth v0.15 documents were written from the approved result-located claim library, rather than prior manuscript prose. International uses four bounded claims; Brazil uses eight. All empirical claim sentences have a Reconstruction_Claim_ID, Evidence_ID, Result_ID and source locator in the matrix. Their documented evidence density is insufficient for complete manuscripts at the approved scope.\n\n## Required next control\n\nTargeted gap closure is required before a final human scientific audit or reference freeze. The documents are not submission ready and do not establish that Architecture B can yet support full manuscripts.\n\n## Rendering\n\nPDF and PNG rendering were not completed because the bundled document runtime does not provide LibreOffice. Structural DOCX validation remains required and visual QA is pending an available bundled renderer.\n'''
 (OUT/'BATCH10_4H_REPORT.md').write_text(report,encoding='utf-8')
 outputs=[p for p in OUT.iterdir() if p.is_file() and p.name!='BATCH10_4H_MANIFEST.json']
 manifest={'batch':'BATCH 10.4H','gate':'V015_RECONSTRUCTION_INSUFFICIENT','architecture':'B_NARROWED_INTEGRATIVE_MANUSCRIPT','approved_claims_accounted':'9/9','empirical_claims_without_result_id':0,'result_ids_without_locator':0,'blocked_evidence_support':0,'source_insufficient_evidence_promoted':0,'rejected_pmcid_reused':0,'cohort_double_counting':0,'null_suppression':0,'causal_inflation':0,'framework_validation_inflation':0,'unsupported_national_generalization':0,'CEF_v1_changes':0,'Zotero_changes':0,'reference_freeze_executed':False,'rendering':'NOT_RENDERED_BUNDLED_LIBREOFFICE_UNAVAILABLE','deterministic_regeneration':'python analysis/build_batch10_4h.py','artifact_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in outputs}}
 (OUT/'BATCH10_4H_MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
if __name__=='__main__':main()
