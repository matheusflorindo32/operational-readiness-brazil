"""Build compact Google Sheets updateCells requests from the verified XLSX."""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

MASTER_SHEET_ID = 1864068772
DASHBOARD_SHEET_ID = 505876673
PROVENANCE_SHEET_ID = 233372420
README_SHEET_ID = 1688208316
GROUPS = ((1, 2), (4, 8), (30, 31), (87, 94), (129, 130), (141, 143))


def sheets_value(value):
    if value is None:
        return {}
    if isinstance(value, bool):
        return {"boolValue": value}
    if isinstance(value, (int, float)):
        return {"numberValue": value}
    if isinstance(value, (date, datetime)):
        serial = (value.date() - date(1899, 12, 30)).days
        return {"numberValue": serial}
    return {"stringValue": str(value)}


def prepared_value(value):
    if isinstance(value, (date, datetime)):
        return (value.date() - date(1899, 12, 30)).days
    return value


def prepare(workbook_path, output_path):
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    master = workbook["Master Evidence"]
    rows = []
    for row in range(4, 1460):
        rows.append(
            [
                [prepared_value(master.cell(row, column).value) for column in range(start + 1, end + 1)]
                for start, end in GROUPS
            ]
        )
    payload = {
        "rows": rows,
        "dashboard_a4": workbook["Dashboard"]["A4"].value,
        "dashboard_n14": workbook["Dashboard"]["N14"].value,
        "provenance_a2": workbook["Search Provenance"]["A2"].value,
        "readme_b10": workbook["README"]["B10"].value,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    workbook.close()


def request_for_group(payload, offset, count, group_index):
    start_column, _ = GROUPS[group_index]
    rows = []
    for grouped_row in payload["rows"][offset : offset + count]:
        rows.append(
            {
                "values": [
                    {"userEnteredValue": sheets_value(value)}
                    for value in grouped_row[group_index]
                ]
            }
        )
    return {
        "updateCells": {
            "start": {
                "sheetId": MASTER_SHEET_ID,
                "rowIndex": 3 + offset,
                "columnIndex": start_column,
            },
            "rows": rows,
            "fields": "userEnteredValue",
        }
    }


def build(payload, offset, count):
    requests = [request_for_group(payload, offset, count, i) for i in range(len(GROUPS))]
    if offset == 0:
        for sheet_id, row, column, value in (
            (DASHBOARD_SHEET_ID, 3, 0, payload["dashboard_a4"]),
            (DASHBOARD_SHEET_ID, 13, 13, payload["dashboard_n14"]),
            (PROVENANCE_SHEET_ID, 1, 0, payload["provenance_a2"]),
            (README_SHEET_ID, 9, 1, payload["readme_b10"]),
        ):
            requests.append(
                {
                    "updateCells": {
                        "start": {"sheetId": sheet_id, "rowIndex": row, "columnIndex": column},
                        "rows": [{"values": [{"userEnteredValue": sheets_value(value)}]}],
                        "fields": "userEnteredValue",
                    }
                }
            )
    return requests


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--prepare", type=Path)
    parser.add_argument("--prepared", type=Path, required=True)
    parser.add_argument("--offset", type=int)
    parser.add_argument("--count", type=int, default=75)
    args = parser.parse_args()
    if args.prepare:
        prepare(args.prepare, args.prepared)
        print(json.dumps({"prepared": str(args.prepared), "rows": 1456}))
        return
    if args.offset is None:
        parser.error("--offset is required unless --prepare is used")
    payload = json.loads(args.prepared.read_text(encoding="utf-8"))
    print(json.dumps(build(payload, args.offset, args.count), ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
